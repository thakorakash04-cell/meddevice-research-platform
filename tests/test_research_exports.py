import io,tempfile
from pathlib import Path
import pytest
from openpyxl import load_workbook
from meddevice_accounts.store import Store,AccessDenied
from research_exports import build_research_export

@pytest.fixture
def accounts():
    with tempfile.TemporaryDirectory() as folder:
        store=Store('sqlite:///'+str(Path(folder)/'db'));store.bootstrap_admin('admin_user','Test-admin-password!123');admin=store.login('admin_user','Test-admin-password!123')
        store.create_viewer(admin,'viewer_user','Test-viewer-password!123')
        store.update_permissions(admin,'viewer_user',{'download':True})
        yield store,admin,store.login('viewer_user','Test-viewer-password!123')

TABLES=[{'id':'T1','name':'FDA 510(k)','headers':['510(k) Number','Device Name','Applicant'],'rows':[[f'K24{i:04}','Laser','=Literal applicant'] for i in range(30)]+[['K999999','Other','Other']]}]

def test_admin_export_includes_all_rows_columns_and_literal_text(accounts):
    store,admin,viewer=accounts
    query={'table':'T1','filters':[{'column':'Device Name','op':'eq','value':'Laser'}],'columns':['510(k) Number'],'limit':20}
    workbook=load_workbook(io.BytesIO(build_research_export(store,admin,TABLES,query)))
    assert workbook['Result table'].max_row==31
    assert workbook['Matching records'].max_row==31
    assert [c.value for c in workbook['Matching records'][1]]==['Source row','510(k) Number','Device Name','Applicant']
    assert workbook['Matching records']['D2'].value=='=Literal applicant' and workbook['Matching records']['D2'].data_type=='s'
    assert any(e['action']=='export' and e['resource']=='Ask AI research matching records' for e in store.access_log(admin))

def test_viewer_with_download_permission_is_denied(accounts):
    store,admin,viewer=accounts
    with pytest.raises(PermissionError,match='administrators'):build_research_export(store,viewer,TABLES,{'table':'T1'})

def test_revoked_admin_and_invalid_tables_are_denied(accounts):
    store,admin,viewer=accounts
    with pytest.raises(ValueError,match='unavailable'):build_research_export(store,admin,TABLES,{'table':'accounts'})
    store.bootstrap_admin('new_admin','Test-new-admin-password!123')
    with pytest.raises(AccessDenied):build_research_export(store,admin,TABLES,{'table':'T1'})

def test_grouped_results_and_matching_records_are_separate(accounts):
    store,admin,viewer=accounts
    query={'table':'T1','group_by':['Device Name'],'metrics':[{'column':'*','op':'count'}]}
    workbook=load_workbook(io.BytesIO(build_research_export(store,admin,TABLES,query)))
    assert workbook['Result table'].max_row==3 and workbook['Matching records'].max_row==32
